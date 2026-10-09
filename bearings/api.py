from rest_framework.viewsets import GenericViewSet
from rest_framework import mixins, status

from bearings.models import *
from bearings.serializers import *

from django.db.models import Q
from rest_framework.decorators import action
from rest_framework.response import Response
from django.db.models import Avg, Max, Min, Sum, Count, F
from rest_framework import serializers
from django.contrib.auth import authenticate, login, logout

from rest_framework.permissions import BasePermission, IsAuthenticated
from django.core.cache import cache

from django.http import HttpResponse
import openpyxl
import io
from datetime import datetime
import pyotp




class OTPRequired(BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and cache.get('otp_good', False))

class ClientsViewset(
    mixins.CreateModelMixin,
    mixins.UpdateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
    ):
    queryset = Client.objects.all()
    serializer_class = ClientSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        
        if not self.request.user.is_superuser:
            qs = qs.filter(user=self.request.user)
        return qs
    
    otp_required_actions = ['update']
    
    def get_permissions(self):
        permissions = super().get_permissions()
        
        if self.action in self.otp_required_actions:
            permissions.append(OTPRequired())
        
        return permissions



class OrdersViewset(
    mixins.CreateModelMixin,
    mixins.UpdateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
    ):
    queryset = Order.objects.all()
    serializer_class = OrderSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        if not self.request.user.is_superuser:
            qs = qs.filter(client__user=self.request.user)
        return qs

class FeedbacksViewset(
    mixins.CreateModelMixin,
    mixins.UpdateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
    ):
    queryset = Feedback.objects.all()
    serializer_class = FeedbackSerializer


class BearingsViewset(
    mixins.CreateModelMixin,
    mixins.UpdateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
    ):
    queryset = Bearing.objects.all()
    serializer_class = BearingSerializer

    class StatsSerializer(serializers.Serializer):
        total_count = serializers.IntegerField()
        total_amount = serializers.IntegerField()
        avg_price = serializers.FloatField()
        max_price = serializers.IntegerField()
        min_price = serializers.IntegerField()
        total_value = serializers.IntegerField()
    
    @action(detail=False, methods=["GET"], url_path="stats")
    def get_stats(self, request, *args, **kwargs):
        stats = Bearing.objects.aggregate(
            total_count=Count("*"),
            total_amount=Sum("ammount"),
            avg_price=Avg("price"),
            max_price=Max("price"),
            min_price=Min("price"),
            total_value=Sum(F("price") * F("ammount"))
        )
        serializer = self.StatsSerializer(instance=stats)
        return Response(serializer.data)
    
    @action(detail=False, methods=["GET"], url_path="export")
    def export(self, request, *args, **kwargs):
        bearings = self.get_queryset()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Подшипники"
        headers = [
            "ID", 
            "Название", 
            "Внутренний диаметр", 
            "Внешний диаметр", 
            "Высота", 
            "Цена", 
            "Количество",
        ]
        
        for col_num, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_num, value=header)
        
        for row_num, bearing in enumerate(bearings, 2):
            ws.cell(row=row_num, column=1, value=bearing.id)
            ws.cell(row=row_num, column=2, value=bearing.name)
            ws.cell(row=row_num, column=3, value=bearing.inner_d)
            ws.cell(row=row_num, column=4, value=bearing.outer_d)
            ws.cell(row=row_num, column=5, value=bearing.height)
            ws.cell(row=row_num, column=6, value=bearing.price)
            ws.cell(row=row_num, column=7, value=bearing.ammount)
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        filename = f"bearings_export_{datetime.now().strftime('%Y-%m-%d_%H-%M')}.xlsx"
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
    
    @action(detail=False, methods=['POST'], url_path='filter')
    def filter_bearings(self, request, *args, **kwargs):
        filters = request.data
        queryset = self.get_queryset()
        
        q_objects = Q()
        
        if filters.get('name'):
            q_objects &= Q(name__icontains=filters['name'])
        
        if filters.get('inner_d_min'):
            q_objects &= Q(inner_d__gte=float(filters['inner_d_min']))
        if filters.get('inner_d_max'):
            q_objects &= Q(inner_d__lte=float(filters['inner_d_max']))
        
        if filters.get('outer_d_min'):
            q_objects &= Q(outer_d__gte=float(filters['outer_d_min']))
        if filters.get('outer_d_max'):
            q_objects &= Q(outer_d__lte=float(filters['outer_d_max']))
        
        if filters.get('height_min'):
            q_objects &= Q(height__gte=float(filters['height_min']))
        if filters.get('height_max'):
            q_objects &= Q(height__lte=float(filters['height_max']))
        
        if filters.get('price_min'):
            q_objects &= Q(price__gte=float(filters['price_min']))
        if filters.get('price_max'):
            q_objects &= Q(price__lte=float(filters['price_max']))
        
        if filters.get('ammount_min'):
            q_objects &= Q(ammount__gte=int(filters['ammount_min']))
        if filters.get('ammount_max'):
            q_objects &= Q(ammount__lte=int(filters['ammount_max']))
        
        filtered_queryset = queryset.filter(q_objects)
        
        serializer = self.get_serializer(filtered_queryset, many=True)
        
        return Response(serializer.data)


class OrdersCompositionsViewset(
    mixins.CreateModelMixin,
    mixins.UpdateModelMixin,
    mixins.RetrieveModelMixin,
    mixins.DestroyModelMixin,
    mixins.ListModelMixin,
    GenericViewSet
    ):
    queryset = OrderComposition.objects.all()
    serializer_class = OrderCompositionSerializer


    class StatsSerializer(serializers.Serializer):
        total_quantity = serializers.IntegerField()
        avg_per_order = serializers.FloatField()
        max_quantity = serializers.IntegerField()
        min_quantity = serializers.IntegerField()
    
    @action(detail=False, methods=["GET"], url_path="stats")
    def get_stats(self, request, *args, **kwargs):
        stats = OrderComposition.objects.aggregate(
            total_quantity=Sum("ammount"),
            avg_per_order=Avg("ammount"),
            max_quantity=Max("ammount"),
            min_quantity=Min("ammount")
        )
        serializer = self.StatsSerializer(instance=stats)
        return Response(serializer.data)

    def get_queryset(self):
        qs = super().get_queryset()
        if not self.request.user.is_superuser:
            qs = qs.filter(order__client__user=self.request.user)
        return qs

    @action(detail=False, methods=['POST'], url_path='filter')
    def filter_order_compositions(self, request, *args, **kwargs):
        filters = request.data
        queryset = self.get_queryset()
        
        q_objects = Q()
        
        if filters.get('order_number'):
            q_objects &= Q(order__number__icontains=filters['order_number'])
        
        if filters.get('bearing_name'):
            q_objects &= Q(bearing__name__icontains=filters['bearing_name'])
        
        if filters.get('ammount_min'):
            q_objects &= Q(ammount__gte=int(filters['ammount_min']))
        if filters.get('ammount_max'):
            q_objects &= Q(ammount__lte=int(filters['ammount_max']))
        
        if filters.get('client_name'):
            q_objects &= Q(order__client__name__icontains=filters['client_name'])
        
        filtered_queryset = queryset.filter(q_objects)

        serializer = self.get_serializer(filtered_queryset, many=True)
        
        return Response(serializer.data)

class UserProfileViewset(GenericViewSet):
    class LoginSerializer(serializers.Serializer):
        username = serializers.CharField()
        password = serializers.CharField()
    
    class RegisterSerializer(serializers.Serializer):
        username = serializers.CharField()
        email = serializers.EmailField()
        password = serializers.CharField()
        phone = serializers.CharField(max_length=20)

    class OTPSerializer(serializers.Serializer):
        key = serializers.CharField()

    @action(url_path="info", detail=False, methods=["GET"])
    def get_user(self, request, *args, **kwargs):
        user = request.user
        data = {"is_authenticated": user.is_authenticated}
        
        if user.is_authenticated:
            data.update({
                "is_superuser": user.is_superuser,
                "name": user.username,
                "is_double": cache.get('otp_good', False) 
        })
            
        return Response(data)     
    
    @action(url_path="login", detail=False, methods=["POST"])
    def login(self, request, *args, **kwargs):
        serializer = self.LoginSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        userdata = serializer.validated_data
        user = authenticate(username = userdata['username'], password = userdata['password'])
        login(request, user)
        if(user is not None):
            return Response(status=status.HTTP_200_OK)
        else:
            return Response(status=status.HTTP_401_UNAUTHORIZED)
            
    
    @action(url_path="logout", detail=False, methods=["POST"])
    def logout(self, request, *args, **kwargs):
        logout(request)
        return Response(status=status.HTTP_200_OK)

    @action(url_path="register", detail=False, methods=["POST"])
    def register(self, request, *args, **kwargs):
        serializer = self.RegisterSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        userdata = serializer.validated_data
        user = User.objects.create_user(
            username=userdata['username'],
            email=userdata['email'],
            password=userdata['password']
        )
        user.save()

        client = Client.objects.create(
            name=userdata['username'],
            phone=userdata['phone'],
            user=user
        )
        client.save()

        user = authenticate(username = userdata['username'], password = userdata['password'])
        login(request, user)

        return Response(status=status.HTTP_200_OK)
    
    @action(url_path="get-totp", methods=['GET'], detail=False, permission_classes=[IsAuthenticated])
    def get_totp(self, *args, **kwargs):
        my_user = Client.objects.filter(user=self.request.user).first()
        
        my_user.totp_key = pyotp.random_base32()
        my_user.save()

        url = pyotp.totp.TOTP(my_user.totp_key).provisioning_uri(
            name=my_user.name, issuer_name="Bearings"
        )

        return Response({
            "url": url
        })

    @action(detail=False, url_path='otp-login', methods=['POST'], serializer_class=OTPSerializer)
    def otp_login(self, *args, **kwargs):
        my_user = Client.objects.filter(user=self.request.user).first()
        totp = pyotp.TOTP(my_user.totp_key)
        
        serializer = self.get_serializer(data=self.request.data)
        serializer.is_valid(raise_exception=True)

        success = False
        if totp.now() == serializer.validated_data['key']:
            cache.set('otp_good', True, 60)
            success = True

        return Response({
            'success': success
        })
    

